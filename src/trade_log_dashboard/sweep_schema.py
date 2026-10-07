"""Parquet sweep schema inspection and automatic dashboard-field mapping."""

from __future__ import annotations

from collections import Counter
import csv
from difflib import SequenceMatcher
from io import BytesIO, StringIO
from pathlib import Path
import re
import tempfile
from typing import Any

import duckdb
from .rejection_report import build_rejection_report


FIELD_DEFINITIONS: tuple[dict[str, Any], ...] = (
    {"key": "status", "label": "Run status", "kind": "text", "required": True,
     "aliases": ("status", "result_status", "run_status")},
    {"key": "net_pnl", "label": "Net P&L", "kind": "numeric", "required": True,
     "aliases": ("net_pnl", "total_pnl", "net_profit", "overall_profit", "strategy_pnl")},
    {"key": "pnl_2025", "label": "2025 P&L", "kind": "numeric", "required": False,
     "aliases": ("pnl_2025", "net_pnl_2025", "profit_2025", "year_2025_pnl")},
    {"key": "pnl_2026", "label": "2026 P&L", "kind": "numeric", "required": False,
     "aliases": ("pnl_2026", "net_pnl_2026", "profit_2026", "year_2026_pnl")},
    {"key": "entry_start", "label": "Entry time", "kind": "time", "required": True,
     "aliases": ("entry_start", "entry_time", "entry", "start_time")},
    {"key": "mtm_dd_2025", "label": "2025 drawdown", "kind": "numeric", "required": False,
     "aliases": ("mtm_dd_2025", "drawdown_2025", "max_drawdown_2025", "dd_2025")},
    {"key": "mtm_dd_2026", "label": "2026 drawdown", "kind": "numeric", "required": False,
     "aliases": ("mtm_dd_2026", "drawdown_2026", "max_drawdown_2026", "dd_2026")},
    {"key": "mtm_drawdown", "label": "Overall drawdown", "kind": "numeric", "required": False,
     "aliases": ("mtm_drawdown", "max_drawdown", "drawdown", "maximum_drawdown")},
    {"key": "completed_trade_count", "label": "Completed trades", "kind": "numeric", "required": False,
     "aliases": ("completed_trade_count", "completed_trades", "trade_count", "total_trades", "num_trades")},
)

PRE_RANKED_MARKER_COLUMNS = {
    "prior_min", "prior_count", "after_min", "after_count",
    "prior_robust_ratio", "after_robust_ratio", "robustness_ratio", "entry_robust_side",
}

_RESULT_COLUMN_NAMES = {
    "status", "win_rate", "worst_day_pnl", "max_consecutive_losses", "final_score", "rank",
    "max_loss", "entry_minutes", "prior_min", "after_min", "drawdown_risk", "max_loss_risk",
    "reentry_burden", "prior_robust_ratio", "after_robust_ratio", "robustness_ratio",
    "entry_robust_side",
}
_RESULT_COLUMN_SUFFIXES = ("_count", "_pnl", "_roi", "_dd", "_drawdown", "_score")


def _is_result_column(name: str) -> bool:
    normalised = _normalise(name)
    return (
        normalised in {_normalise(item) for item in _RESULT_COLUMN_NAMES}
        or normalised.endswith(tuple(_normalise(item) for item in _RESULT_COLUMN_SUFFIXES))
        or bool(re.fullmatch(r"(?:pnl|roi|mtmdd|drawdown)\d{4}(?:pct)?", normalised))
    )


def _parameter_filter_clauses(parameter_filters: Any, columns: dict[str, dict[str, Any]], identifier) -> tuple[list[str], list[Any]]:
    """Build exact-value or numeric-range filters; values in one column are ORed."""
    if parameter_filters in (None, ""):
        return [], []
    items = parameter_filters if isinstance(parameter_filters, list) else [
        {"column": name, "mode": "value", "value": value} for name, value in parameter_filters.items()
    ] if isinstance(parameter_filters, dict) else None
    if items is None:
        raise ValueError("Combination filters must be a list of selected values or ranges.")
    exact_values: dict[str, list[str]] = {}
    numeric_ranges: dict[str, list[tuple[float, float, bool]]] = {}
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("Each combination filter must be an object.")
        name = str(item.get("column", "")).strip()
        mode = str(item.get("mode", "value")).strip().lower()
        if not name:
            continue
        if name not in columns:
            raise ValueError(f'Combination filter column "{name}" is not present in this Parquet file.')
        column = columns[name]
        if mode in {"range", "ranges"}:
            if column["category"] != "numeric":
                raise ValueError(f'Range filters are only available for numeric column "{name}".')
            raw_ranges = item.get("ranges") if mode == "ranges" else [item]
            if not isinstance(raw_ranges, list) or not raw_ranges:
                raise ValueError(f'Range filter for "{name}" is invalid.')
            for raw_range in raw_ranges:
                try:
                    minimum, maximum = float(raw_range["min"]), float(raw_range["max"])
                except (KeyError, TypeError, ValueError) as exc:
                    raise ValueError(f'Range filter for "{name}" is invalid.') from exc
                if minimum > maximum:
                    raise ValueError(f'Range filter for "{name}" must have its minimum before its maximum.')
                numeric_ranges.setdefault(name, []).append((minimum, maximum, bool(raw_range.get("include_max", False))))
            continue
        if mode not in {"value", "values"}:
            raise ValueError(f'Filter mode for "{name}" must be value, values, range, or ranges.')
        raw_values = item.get("values") if mode == "values" else item.get("value", "")
        selected = raw_values if isinstance(raw_values, list) else str(raw_values).split(",")
        exact_values.setdefault(name, []).extend(str(value).strip() for value in selected if str(value).strip())

    clauses, values = [], []
    for name, selected in exact_values.items():
        selected = list(dict.fromkeys(selected))
        if not selected:
            continue
        if columns[name]["category"] == "numeric":
            try:
                numeric = [float(value.replace("%", "")) for value in selected]
            except ValueError as exc:
                raise ValueError(f'Filter for "{name}" must contain numeric values.') from exc
            clauses.append(f"TRY_CAST({identifier(name)} AS DOUBLE) IN ({', '.join('?' for _ in numeric)})")
            values.extend(numeric)
        else:
            clauses.append(f"lower(CAST({identifier(name)} AS VARCHAR)) IN ({', '.join('lower(?)' for _ in selected)})")
            values.extend(selected)
    for name, ranges in numeric_ranges.items():
        range_clauses = []
        for minimum, maximum, include_max in ranges:
            operator = "<=" if include_max else "<"
            range_clauses.append(
                f"(TRY_CAST({identifier(name)} AS DOUBLE) >= ? AND TRY_CAST({identifier(name)} AS DOUBLE) {operator} ?)"
            )
            values.extend([minimum, maximum])
        if range_clauses:
            clauses.append("(" + " OR ".join(range_clauses) + ")")
    return clauses, values


def _export_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Flatten ranked combinations while retaining every parameter and score component."""
    if not rows:
        raise ValueError("Rank at least one eligible combination before exporting.")
    parameter_names = sorted({str(name) for row in rows for name in (row.get("parameters") or {})})
    component_count = max(len(row.get("ranking_components") or []) for row in rows)
    flattened = []
    for source in rows:
        row = {
            "rank": source.get("rank"), "final_score": source.get("final_score"),
            "net_pnl": source.get("net_pnl"), "pnl_2025": source.get("pnl_2025"),
            "pnl_2026": source.get("pnl_2026"), "ranking_drawdown": source.get("ranking_drawdown"),
            "entry_start": source.get("entry_start"), "entry_robustness": source.get("entry_robustness"),
        }
        parameters = source.get("parameters") or {}
        row.update({f"parameter_{name}": parameters.get(name) for name in parameter_names})
        components = source.get("ranking_components") or []
        for index in range(component_count):
            prefix = f"ranking_component_{index + 1}_"
            component = components[index] if index < len(components) else {}
            row.update({
                prefix + "metric": component.get("label"),
                prefix + "column": component.get("column"),
                prefix + "direction": component.get("direction"),
                prefix + "weight": component.get("weight"),
                prefix + "value": component.get("value"),
                prefix + "percentile_score": component.get("percentile_score"),
                prefix + "weighted_contribution": component.get("weighted_contribution"),
            })
        flattened.append(row)
    return flattened


def export_ranked_sweep(rows: list[dict[str, Any]], export_format: str) -> bytes:
    """Create a CSV, Parquet, or XLSX export for the current ranked results."""
    records = _export_rows(rows)
    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=list(records[0]), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
        return output.getvalue().encode("utf-8-sig")
    if export_format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover - deployment dependency guard
            raise ValueError("Excel export requires openpyxl to be installed.") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Ranked strategies"
        headers = list(records[0])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for record in records:
            sheet.append([record.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            longest = max(len(str(header)), *(len(str(record.get(header) or "")) for record in records))
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 35)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    if export_format == "parquet":
        try:
            import pandas as pd
        except ImportError as exc:  # pragma: no cover - deployment dependency guard
            raise ValueError("Parquet export requires pandas to be installed.") from exc
        with tempfile.TemporaryDirectory(prefix="ranked-sweep-export-") as temporary:
            path = Path(temporary, "ranked_strategies.parquet")
            connection = duckdb.connect()
            try:
                connection.register("export_rows", pd.DataFrame(records))
                connection.execute("COPY export_rows TO ? (FORMAT PARQUET)", [str(path)])
            finally:
                connection.close()
            return path.read_bytes()
    raise ValueError("Export format must be csv, parquet, or xlsx.")


def _normalise(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _category(duckdb_type: str) -> str:
    upper = duckdb_type.upper()
    if any(token in upper for token in ("INT", "FLOAT", "DOUBLE", "DECIMAL", "HUGEINT", "UBIGINT")):
        return "numeric"
    if any(token in upper for token in ("DATE", "TIME", "TIMESTAMP")):
        return "time"
    if any(token in upper for token in ("CHAR", "VARCHAR", "STRING", "ENUM")):
        return "text"
    if upper == "BOOLEAN":
        return "boolean"
    return "other"


def _compatible(expected: str, actual: str) -> bool:
    if expected == "numeric":
        return actual == "numeric"
    if expected == "time":
        return actual in {"time", "text", "numeric"}
    if expected == "text":
        return actual in {"text", "boolean"}
    return True


def _suggest(field: dict[str, Any], columns: list[dict[str, Any]], used: set[str]) -> tuple[str | None, str | None]:
    aliases = tuple(_normalise(alias) for alias in field["aliases"])
    for column in columns:
        if column["name"] not in used and _normalise(column["name"]) in aliases:
            return column["name"], "exact"

    best_name, best_score = None, 0.0
    for column in columns:
        if column["name"] in used:
            continue
        candidate = _normalise(column["name"])
        score = max(SequenceMatcher(None, candidate, alias).ratio() for alias in aliases)
        if score > best_score:
            best_name, best_score = column["name"], score
    return (best_name, "similar") if best_score >= 0.76 else (None, None)


def _filter_option_label(value: Any, *, numeric: bool = False) -> str:
    if value is None:
        return "—"
    if numeric:
        number = float(value)
        return str(int(number)) if number.is_integer() else f"{number:.6g}"
    return str(value)


def _build_filter_options(path: Path, columns: list[dict[str, Any]], mapped_sources: set[str]) -> dict[str, dict[str, Any]]:
    """Build exact-value or ten-bucket options for unmapped strategy parameters."""
    options: dict[str, dict[str, Any]] = {}
    candidates = [column for column in columns if column["name"] not in mapped_sources
                  and not _is_result_column(column["name"])
                  and column["category"] in {"numeric", "text", "time", "boolean"}]
    connection = duckdb.connect()

    def identifier(name: str) -> str:
        return '"' + name.replace('"', '""') + '"'

    try:
        for column in candidates:
            name, quoted = column["name"], identifier(column["name"])
            distinct = [row[0] for row in connection.execute(
                f"SELECT DISTINCT CAST({quoted} AS VARCHAR) FROM read_parquet(?) WHERE {quoted} IS NOT NULL ORDER BY 1 LIMIT 101",
                [str(path)],
            ).fetchall()]
            if len(distinct) <= 100:
                numeric = column["category"] == "numeric"
                if numeric:
                    distinct.sort(key=lambda value: float(value))
                options[name] = {"mode": "values", "category": column["category"], "values": [
                    {"label": _filter_option_label(value, numeric=numeric), "value": str(value)} for value in distinct
                ]}
                continue
            if column["category"] != "numeric":
                options[name] = {"mode": "search", "category": column["category"], "values": []}
                continue
            minimum, maximum = connection.execute(
                f"SELECT min(TRY_CAST({quoted} AS DOUBLE)), max(TRY_CAST({quoted} AS DOUBLE)) FROM read_parquet(?) WHERE {quoted} IS NOT NULL",
                [str(path)],
            ).fetchone()
            if minimum is None or maximum is None or float(minimum) == float(maximum):
                options[name] = {"mode": "values", "category": column["category"], "values":
                    [{"label": _filter_option_label(minimum, numeric=True), "value": str(minimum)}] if minimum is not None else []}
                continue
            minimum, maximum = float(minimum), float(maximum)
            width = (maximum - minimum) / 10.0
            ranges = []
            for index in range(10):
                lower = minimum + width * index
                upper = maximum if index == 9 else minimum + width * (index + 1)
                ranges.append({"label": f"{_filter_option_label(lower, numeric=True)} – {_filter_option_label(upper, numeric=True)}{' (inclusive)' if index == 9 else ''}",
                               "min": lower, "max": upper, "include_max": index == 9})
            options[name] = {"mode": "ranges", "category": column["category"], "values": ranges}
    finally:
        connection.close()
    return options


def inspect_sweep_parquet(path: Path, *, include_filter_options: bool = True) -> dict[str, Any]:
    """Read Parquet metadata and return mappings plus a validation report."""
    path = Path(path)
    connection = duckdb.connect()
    try:
        schema_rows = connection.execute(
            """SELECT name, COALESCE(duckdb_type, type, 'UNKNOWN') AS data_type,
                      repetition_type, column_id
                 FROM parquet_schema(?)
                WHERE num_children IS NULL
                ORDER BY column_id""",
            [str(path)],
        ).fetchall()
        row_count = sum(
            int(row[0]) for row in connection.execute(
                "SELECT num_rows FROM parquet_file_metadata(?)", [str(path)]
            ).fetchall()
        )
    except duckdb.Error as exc:
        raise ValueError(f"The selected file is not a readable Parquet sweep: {exc}") from exc
    finally:
        connection.close()

    columns = [
        {
            "name": str(name),
            "type": str(data_type),
            "category": _category(str(data_type)),
            "nullable": str(repetition).upper() != "REQUIRED",
        }
        for name, data_type, repetition, _column_id in schema_rows
    ]
    if not columns:
        raise ValueError("The Parquet file has no leaf columns.")

    duplicates = sorted(name for name, count in Counter(column["name"] for column in columns).items() if count > 1)
    used: set[str] = set()
    suggestions: dict[str, tuple[str | None, str | None]] = {}
    # Reserve exact aliases first so a fuzzy optional match cannot take a
    # column that is an exact match for a later field.
    for definition in FIELD_DEFINITIONS:
        aliases = {_normalise(alias) for alias in definition["aliases"]}
        source = next((column["name"] for column in columns
                       if column["name"] not in used and _normalise(column["name"]) in aliases), None)
        if source:
            suggestions[definition["key"]] = (source, "exact")
            used.add(source)
    for definition in FIELD_DEFINITIONS:
        if definition["key"] not in suggestions:
            source, confidence = _suggest(definition, columns, used)
            suggestions[definition["key"]] = (source, confidence)
            if source:
                used.add(source)

    fields = []
    for definition in FIELD_DEFINITIONS:
        source, confidence = suggestions[definition["key"]]
        fields.append({
            "key": definition["key"],
            "label": definition["label"],
            "kind": definition["kind"],
            "required": definition["required"],
            "suggested_source": source,
            "confidence": confidence,
        })

    by_name = {column["name"]: column for column in columns}
    errors: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []
    pre_ranked_markers = sorted(
        column["name"] for column in columns
        if _normalise(column["name"]) in {_normalise(name) for name in PRE_RANKED_MARKER_COLUMNS}
    )
    if len(pre_ranked_markers) >= 2:
        warnings.append({
            "code": "pre_ranked_output",
            "field": "raw_sweep",
            "message": (
                "This file includes prior ranking or robustness columns. It can still be ranked as base combinations, "
                "but entry-time robustness can only be confirmed when matching nearby entry variants are present."
            ),
        })
    for name in duplicates:
        errors.append({"code": "duplicate_column", "field": name,
                       "message": f'Column "{name}" appears more than once.'})
    for field in fields:
        source = field["suggested_source"]
        if field["required"] and not source:
            errors.append({"code": "missing_mapping", "field": field["key"],
                           "message": f'Map a Parquet column to {field["label"]}.'})
        elif source and not _compatible(field["kind"], by_name[source]["category"]):
            errors.append({"code": "invalid_type", "field": field["key"],
                           "message": f'{field["label"]} expects {field["kind"]} data; "{source}" is {by_name[source]["type"]}.'})

    mapping = {field["key"]: field["suggested_source"] for field in fields}
    yearly_drawdown = bool(mapping["mtm_dd_2025"] and mapping["mtm_dd_2026"])
    if not mapping["mtm_drawdown"] and not yearly_drawdown:
        errors.append({"code": "missing_drawdown", "field": "drawdown",
                       "message": "Map overall drawdown, or map both 2025 and 2026 drawdown."})
    if not mapping["completed_trade_count"]:
        warnings.append({"code": "optional_mapping", "field": "completed_trade_count",
                         "message": "Completed trades is not mapped; trade-count filtering will be unavailable."})

    result = {
        "file": {"name": path.name, "size_bytes": path.stat().st_size, "row_count": row_count},
        "columns": columns,
        "fields": fields,
        "validation": {
            "ready": not errors,
            "errors": errors,
            "warnings": warnings,
            "duplicate_columns": duplicates,
        },
    }
    if include_filter_options:
        result["filter_options"] = _build_filter_options(path, columns, {source for source in mapping.values() if source})
    return result


def preview_sweep_filters(
    path: Path,
    mapping: dict[str, str],
    filters: dict[str, Any],
) -> dict[str, Any]:
    """Count rows meeting dashboard filter criteria without loading the sweep into memory."""
    report = inspect_sweep_parquet(path, include_filter_options=False)
    columns = {column["name"]: column for column in report["columns"]}
    fields = {field["key"]: field for field in report["fields"]}

    def source_for(key: str, *, required: bool = False) -> str | None:
        source = mapping.get(key)
        if not source:
            if required:
                raise ValueError(f"Map a Parquet column to {fields[key]['label']} before filtering.")
            return None
        if source not in columns:
            raise ValueError(f'Mapped column "{source}" is not present in this Parquet file.')
        if not _compatible(fields[key]["kind"], columns[source]["category"]):
            raise ValueError(
                f'{fields[key]["label"]} expects {fields[key]["kind"]} data; '
                f'"{source}" is {columns[source]["type"]}.'
            )
        return source

    def identifier(name: str) -> str:
        return '"' + name.replace('"', '""') + '"'

    def number(value: Any, label: str) -> float | None:
        if value in (None, ""):
            return None
        try:
            return float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{label} must be a number.") from exc

    status_column = source_for("status", required=True)
    pnl_column = source_for("net_pnl", required=True)
    entry_column = source_for("entry_start", required=True)
    yearly_2025 = source_for("mtm_dd_2025")
    yearly_2026 = source_for("mtm_dd_2026")
    overall_drawdown = source_for("mtm_drawdown")
    trade_count = source_for("completed_trade_count")
    if not overall_drawdown and not (yearly_2025 and yearly_2026):
        raise ValueError("Map overall drawdown, or map both 2025 and 2026 drawdown before filtering.")

    clauses: list[str] = []
    values: list[Any] = [str(path)]
    selected_status = str(filters.get("status", "")).strip()
    if selected_status:
        clauses.append(f"lower(CAST({identifier(status_column)} AS VARCHAR)) = lower(?)")
        values.append(selected_status)

    minimum_pnl = number(filters.get("minimum_pnl"), "Minimum P&L")
    if minimum_pnl is not None:
        clauses.append(f"CAST({identifier(pnl_column)} AS DOUBLE) > ?")
        values.append(minimum_pnl)
    maximum_pnl = number(filters.get("maximum_pnl"), "Maximum P&L")
    if maximum_pnl is not None:
        clauses.append(f"CAST({identifier(pnl_column)} AS DOUBLE) <= ?")
        values.append(maximum_pnl)
    if minimum_pnl is not None and maximum_pnl is not None and maximum_pnl <= minimum_pnl:
        raise ValueError("Maximum P&L must be greater than minimum P&L.")

    latest_entry = str(filters.get("latest_entry", "")).strip()
    if latest_entry:
        if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", latest_entry):
            raise ValueError("Latest entry time must use HH:MM in 24-hour time.")
        if columns[entry_column]["category"] == "numeric":
            hour, minute = (int(part) for part in latest_entry.split(":"))
            clauses.append(f"CAST({identifier(entry_column)} AS DOUBLE) <= ?")
            values.append(hour * 60 + minute)
        else:
            clauses.append(f"TRY_CAST(CAST({identifier(entry_column)} AS VARCHAR) AS TIME) <= CAST(? AS TIME)")
            values.append(latest_entry)

    maximum_drawdown = number(filters.get("maximum_drawdown"), "Maximum drawdown")
    if maximum_drawdown is not None:
        if maximum_drawdown < 0:
            raise ValueError("Maximum drawdown must be zero or greater.")
        if overall_drawdown:
            drawdown_expression = f"abs(CAST({identifier(overall_drawdown)} AS DOUBLE))"
        else:
            drawdown_expression = (
                f"greatest(abs(CAST({identifier(yearly_2025)} AS DOUBLE)), "
                f"abs(CAST({identifier(yearly_2026)} AS DOUBLE)))"
            )
        clauses.append(drawdown_expression + " <= ?")
        values.append(maximum_drawdown)

    minimum_trades = number(filters.get("minimum_trades"), "Minimum trade count")
    if minimum_trades is not None:
        if minimum_trades < 0 or not minimum_trades.is_integer():
            raise ValueError("Minimum trade count must be a whole number greater than or equal to zero.")
        if not trade_count:
            raise ValueError("Map a completed-trades column before applying a trade-count filter.")
        clauses.append(f"CAST({identifier(trade_count)} AS DOUBLE) >= ?")
        values.append(minimum_trades)
    parameter_clauses, parameter_values = _parameter_filter_clauses(filters.get("parameter_filters"), columns, identifier)
    clauses.extend(parameter_clauses)
    values.extend(parameter_values)

    query = "SELECT count(*) FROM read_parquet(?)"
    if clauses:
        query += " WHERE " + " AND ".join(clauses)
    connection = duckdb.connect()
    try:
        matching_rows = int(connection.execute(query, values).fetchone()[0])
    except duckdb.Error as exc:
        raise ValueError(f"The selected filters could not be applied: {exc}") from exc
    finally:
        connection.close()
    return {
        "matching_rows": matching_rows,
        "total_rows": report["file"]["row_count"],
        "applied_filters": {
            "status": selected_status or None,
            "minimum_pnl": minimum_pnl,
            "maximum_pnl": maximum_pnl,
            "maximum_drawdown": maximum_drawdown,
            "latest_entry": latest_entry or None,
            "minimum_trades": int(minimum_trades) if minimum_trades is not None else None,
            "parameter_filters": filters.get("parameter_filters") or [],
        },
    }


def rank_sweep_parquet(
    path: Path,
    mapping: dict[str, str],
    filters: dict[str, Any],
    *,
    top_n: int = 20,
    ranking: list[dict[str, Any]] | None = None,
    require_robustness: bool = True,
) -> dict[str, Any]:
    """Apply eligibility and ranking, optionally requiring entry-time robustness."""
    if not 1 <= top_n <= 100:
        raise ValueError("Top-N must be between 1 and 100.")
    report = inspect_sweep_parquet(path, include_filter_options=False)
    columns = {column["name"]: column for column in report["columns"]}
    fields = {field["key"]: field for field in report["fields"]}

    def source_for(key: str, *, required: bool = False) -> str | None:
        source = mapping.get(key)
        if not source:
            if required:
                raise ValueError(f"Map a Parquet column to {fields[key]['label']} before ranking.")
            return None
        if source not in columns or not _compatible(fields[key]["kind"], columns[source]["category"]):
            raise ValueError(f"The mapping for {fields[key]['label']} is invalid.")
        return source

    def identifier(name: str) -> str:
        return '"' + name.replace('"', '""') + '"'

    def number(value: Any, label: str) -> float | None:
        if value in (None, ""):
            return None
        try:
            return float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{label} must be a number.") from exc

    status = source_for("status", required=True)
    net_pnl = source_for("net_pnl", required=True)
    pnl_2025 = source_for("pnl_2025")
    pnl_2026 = source_for("pnl_2026")
    entry = source_for("entry_start", required=True)
    dd_2025 = source_for("mtm_dd_2025")
    dd_2026 = source_for("mtm_dd_2026")
    overall_dd = source_for("mtm_drawdown")
    trade_count = source_for("completed_trade_count")
    if not pnl_2025 and not pnl_2026:
        raise ValueError("Map at least one yearly P&L column before ranking.")
    if not overall_dd and not (dd_2025 and dd_2026):
        raise ValueError("Map overall drawdown, or map both 2025 and 2026 drawdown before ranking.")

    if ranking is not None:
        raw_criteria = ranking
    elif pnl_2025 and pnl_2026:
        raw_criteria = [
            {"column": pnl_2025, "weight": 30, "direction": "higher", "label": "2025 P&L"},
            {"column": pnl_2026, "weight": 30, "direction": "higher", "label": "2026 P&L"},
            {"column": "__ranking_drawdown", "weight": 40, "direction": "lower", "label": "Ranking drawdown"},
        ]
    else:
        yearly_pnl = pnl_2025 or pnl_2026
        raw_criteria = [
            {"column": yearly_pnl, "weight": 60, "direction": "higher", "label": "Yearly P&L"},
            {"column": "__ranking_drawdown", "weight": 40, "direction": "lower", "label": "Ranking drawdown"},
        ]
    if not isinstance(raw_criteria, list) or not raw_criteria or len(raw_criteria) > 6:
        raise ValueError("Choose between one and six ranking criteria.")
    criteria: list[dict[str, Any]] = []
    used_criteria: set[str] = set()
    for index, criterion in enumerate(raw_criteria):
        if not isinstance(criterion, dict):
            raise ValueError("Each ranking criterion must be an object.")
        column = criterion.get("column")
        direction = criterion.get("direction")
        if not isinstance(column, str) or not column:
            raise ValueError("Choose a column for every ranking criterion.")
        if direction not in {"higher", "lower"}:
            raise ValueError("Ranking direction must be either higher or lower.")
        weight = number(criterion.get("weight"), "Ranking weight")
        if weight is None or weight <= 0:
            raise ValueError("Every ranking weight must be greater than zero.")
        if column in used_criteria:
            raise ValueError("Use each ranking column only once.")
        used_criteria.add(column)
        if column == "__ranking_drawdown":
            source, label = "ranking_drawdown", "Ranking drawdown"
        elif column in columns and columns[column]["category"] == "numeric":
            source, label = column, str(criterion.get("label") or column)
        else:
            raise ValueError(f'Ranking column "{column}" must be a numeric Parquet column.')
        criteria.append({"column": column, "source": source, "label": label, "weight": weight, "direction": direction, "index": index})
    if abs(sum(item["weight"] for item in criteria) - 100) > 0.001:
        raise ValueError("Ranking weights must total 100%.")

    # Sweep result fields must not participate in the “all other settings are
    # unchanged” key. Parameters remain intact, including exit_time.
    mapped_sources = {source for source in mapping.values() if source}
    result_suffixes = ("_count", "_pnl", "_roi", "_dd", "_drawdown", "_score")
    result_names = {
        "status", "win_rate", "worst_day_pnl", "max_consecutive_losses", "final_score", "rank",
        "max_loss", "entry_minutes", "prior_min", "after_min", "drawdown_risk", "max_loss_risk",
        "reentry_burden", "prior_robust_ratio", "after_robust_ratio", "robustness_ratio",
        "entry_robust_side",
    }
    def is_result_column(name: str) -> bool:
        normalised = _normalise(name)
        return (
            normalised in {_normalise(item) for item in result_names}
            or normalised.endswith(tuple(_normalise(item) for item in result_suffixes))
            or bool(re.fullmatch(r"(?:pnl|roi|mtmdd|drawdown)\d{4}(?:pct)?", normalised))
        )

    parameter_columns = [
        column["name"] for column in report["columns"]
        if column["name"] not in mapped_sources and not is_result_column(column["name"])
    ]
    if not parameter_columns:
        raise ValueError("No strategy-parameter columns were detected for entry-time robustness.")
    # Some sweep producers write cutoff_time as a duplicate of entry_start.
    # Keep it in the output, but omit it from the fixed-settings key so that
    # nearby entry-time variants can be compared.
    robustness_key_columns = [
        column for column in parameter_columns
        if _normalise(column) not in {"cutofftime"}
    ]

    minimum_pnl = number(filters.get("minimum_pnl"), "Minimum P&L")
    maximum_pnl = number(filters.get("maximum_pnl"), "Maximum P&L")
    maximum_drawdown = number(filters.get("maximum_drawdown"), "Maximum drawdown")
    minimum_trades = number(filters.get("minimum_trades"), "Minimum trade count")
    selected_status = str(filters.get("status", "")).strip()
    latest_entry = str(filters.get("latest_entry", "")).strip()
    if minimum_pnl is not None and maximum_pnl is not None and maximum_pnl <= minimum_pnl:
        raise ValueError("Maximum P&L must be greater than minimum P&L.")
    if maximum_drawdown is not None and maximum_drawdown < 0:
        raise ValueError("Maximum drawdown must be zero or greater.")
    if minimum_trades is not None and (minimum_trades < 0 or not minimum_trades.is_integer()):
        raise ValueError("Minimum trade count must be a whole number greater than or equal to zero.")
    if minimum_trades is not None and not trade_count:
        raise ValueError("Map a completed-trades column before applying a trade-count filter.")
    if latest_entry and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", latest_entry):
        raise ValueError("Latest entry time must use HH:MM in 24-hour time.")

    entry_expression = (
        f"CAST({identifier(entry)} AS DOUBLE)"
        if columns[entry]["category"] == "numeric"
        else f"extract(hour FROM TRY_CAST(CAST({identifier(entry)} AS VARCHAR) AS TIME)) * 60 + extract(minute FROM TRY_CAST(CAST({identifier(entry)} AS VARCHAR) AS TIME))"
    )
    if overall_dd:
        drawdown_expression = f"abs(CAST({identifier(overall_dd)} AS DOUBLE))"
    else:
        drawdown_expression = f"greatest(abs(CAST({identifier(dd_2025)} AS DOUBLE)), abs(CAST({identifier(dd_2026)} AS DOUBLE)))"
    parameter_key = "hash(" + ", ".join(identifier(column) for column in robustness_key_columns) + ")"
    parameter_select = ", ".join(f"any_value({identifier(column)}) AS {identifier(column)}" for column in parameter_columns)
    output_parameters = ", ".join(identifier(column) for column in parameter_columns)
    custom_metric_select = ", ".join(
        f"any_value(CAST({identifier(item['source'])} AS DOUBLE)) AS custom_metric_{item['index']}"
        for item in criteria if item["source"] != "ranking_drawdown"
    )
    custom_metric_sql = ", " + custom_metric_select if custom_metric_select else ""

    base_clauses = [
        f"lower(CAST({identifier(status)} AS VARCHAR)) = 'succeeded'",
        f"CAST({identifier(net_pnl)} AS DOUBLE) > 0",
        drawdown_expression + " IS NOT NULL",
    ]
    if pnl_2025:
        base_clauses.append(f"CAST({identifier(pnl_2025)} AS DOUBLE) IS NOT NULL")
    if pnl_2026:
        base_clauses.append(f"CAST({identifier(pnl_2026)} AS DOUBLE) IS NOT NULL")
    for item in criteria:
        if item["source"] != "ranking_drawdown":
            base_clauses.append(f"CAST({identifier(item['source'])} AS DOUBLE) IS NOT NULL")
    values: list[Any] = [str(path)]
    if selected_status:
        base_clauses.append(f"lower(CAST({identifier(status)} AS VARCHAR)) = lower(?)")
        values.append(selected_status)
    if minimum_pnl is not None:
        base_clauses.append(f"CAST({identifier(net_pnl)} AS DOUBLE) > ?")
        values.append(minimum_pnl)
    if maximum_pnl is not None:
        base_clauses.append(f"CAST({identifier(net_pnl)} AS DOUBLE) <= ?")
        values.append(maximum_pnl)
    if latest_entry:
        if columns[entry]["category"] == "numeric":
            hour, minute = (int(part) for part in latest_entry.split(":"))
            base_clauses.append("entry_minutes <= ?")
            values.append(hour * 60 + minute)
        else:
            base_clauses.append("entry_minutes <= ?")
            hour, minute = (int(part) for part in latest_entry.split(":"))
            values.append(hour * 60 + minute)
    if maximum_drawdown is not None:
        base_clauses.append("ranking_drawdown <= ?")
        values.append(maximum_drawdown)
    if minimum_trades is not None:
        base_clauses.append(f"CAST({identifier(trade_count)} AS DOUBLE) >= ?")
        values.append(minimum_trades)
    parameter_clauses, parameter_values = _parameter_filter_clauses(filters.get("parameter_filters"), columns, identifier)
    base_clauses.extend(parameter_clauses)
    values.extend(parameter_values)

    # Adapt to the sweep's actual entry-time grid. A two-minute sweep tests
    # ten variants per side; a ten-minute sweep tests the available ±10/±20
    # variants instead of failing solely because two-minute rows are absent.
    robustness_window_minutes = 20

    score_selects, weighted_scores = [], []
    for item in criteria:
        metric = "ranking_drawdown" if item["source"] == "ranking_drawdown" else f"custom_metric_{item['index']}"
        percentile = f"percent_rank() OVER (ORDER BY {metric})"
        score = percentile if item["direction"] == "higher" else f"1 - ({percentile})"
        score_selects.append(f"{score} AS criterion_score_{item['index']}")
        weighted_scores.append(f"{item['weight'] / 100.0} * criterion_score_{item['index']}")
    final_score_expression = " + ".join(weighted_scores)
    component_selects = []
    for item in criteria:
        metric = "ranking_drawdown" if item["source"] == "ranking_drawdown" else f"custom_metric_{item['index']}"
        component_selects.append(
            f"{metric} AS criterion_value_{item['index']}, criterion_score_{item['index']}"
        )
    component_select = ", ".join(component_selects)

    robustness_filter = "WHERE before_pass OR after_pass" if require_robustness else ""
    query = f"""
        WITH source AS (
          SELECT *, {parameter_key} AS parameter_key,
                 {entry_expression} AS entry_minutes,
                 {drawdown_expression} AS ranking_drawdown
            FROM read_parquet(?)
        ),
        time_pnl AS (
          SELECT parameter_key, entry_minutes,
                 max(CASE WHEN lower(CAST({identifier(status)} AS VARCHAR)) = 'succeeded'
                          THEN CAST({identifier(net_pnl)} AS DOUBLE) END) AS shifted_pnl
            FROM source
           WHERE entry_minutes IS NOT NULL
           GROUP BY parameter_key, entry_minutes
        ),
        base AS (
          SELECT parameter_key, entry_minutes,
                 any_value(CAST({identifier(entry)} AS VARCHAR)) AS entry_start,
                 max(CAST({identifier(net_pnl)} AS DOUBLE)) AS net_pnl,
                 {f'any_value(CAST({identifier(pnl_2025)} AS DOUBLE))' if pnl_2025 else 'NULL::DOUBLE'} AS pnl_2025,
                 {f'any_value(CAST({identifier(pnl_2026)} AS DOUBLE))' if pnl_2026 else 'NULL::DOUBLE'} AS pnl_2026,
                 any_value(ranking_drawdown) AS ranking_drawdown,
                 {parameter_select}{custom_metric_sql}
            FROM source
           WHERE {' AND '.join(base_clauses)}
           GROUP BY parameter_key, entry_minutes
        ),
        robustness_stats AS (
          SELECT b.*,
                 count(variant.entry_minutes) FILTER (WHERE variant.entry_minutes < b.entry_minutes) AS before_variant_count,
                 count(variant.entry_minutes) FILTER (WHERE variant.entry_minutes > b.entry_minutes) AS after_variant_count,
                 bool_and(coalesce(variant.shifted_pnl >= 0.70 * b.net_pnl, false))
                   FILTER (WHERE variant.entry_minutes < b.entry_minutes) AS before_all_pass,
                 bool_and(coalesce(variant.shifted_pnl >= 0.70 * b.net_pnl, false))
                   FILTER (WHERE variant.entry_minutes > b.entry_minutes) AS after_all_pass
            FROM base b
            LEFT JOIN time_pnl variant ON variant.parameter_key = b.parameter_key
                                  AND variant.entry_minutes <> b.entry_minutes
                                  AND variant.entry_minutes BETWEEN b.entry_minutes - {robustness_window_minutes}
                                                               AND b.entry_minutes + {robustness_window_minutes}
           GROUP BY ALL
        ),
        robustness AS (
          SELECT *,
                 before_variant_count > 0 AND coalesce(before_all_pass, false) AS before_pass,
                 after_variant_count > 0 AND coalesce(after_all_pass, false) AS after_pass
            FROM robustness_stats
        ),
        eligible AS (
          SELECT *, CASE WHEN before_pass AND after_pass THEN 'both'
                         WHEN before_pass THEN 'before'
                         WHEN after_pass THEN 'after'
                         ELSE 'not_confirmed' END AS entry_robustness
            FROM robustness
           {robustness_filter}
        ),
        scored AS (
          SELECT *,
                 {', '.join(score_selects)}
            FROM eligible
        ),
        score_values AS (
          SELECT *, {final_score_expression} AS final_score
            FROM scored
        ),
        distribution AS (
          SELECT count(*) AS eligible_count,
                 count(*) FILTER (WHERE entry_robustness = 'before') AS before_count,
                 count(*) FILTER (WHERE entry_robustness = 'after') AS after_count,
                 count(*) FILTER (WHERE entry_robustness = 'both') AS both_count,
                 count(*) FILTER (WHERE entry_robustness = 'not_confirmed') AS not_confirmed_count
            FROM score_values
        ),
        top_rows AS (
          SELECT * FROM score_values
           ORDER BY final_score DESC, net_pnl DESC, ranking_drawdown ASC,
                    parameter_key ASC, entry_minutes ASC
           LIMIT {int(top_n)}
        )
        SELECT row_number() OVER (ORDER BY final_score DESC, net_pnl DESC, ranking_drawdown ASC,
                                  parameter_key ASC, entry_minutes ASC) AS rank,
               final_score, net_pnl, pnl_2025, pnl_2026, ranking_drawdown, entry_start, entry_robustness, before_variant_count, after_variant_count,
               eligible_count, before_count, after_count, both_count, not_confirmed_count, {component_select}, {output_parameters}
          FROM top_rows CROSS JOIN distribution
         ORDER BY rank
    """
    connection = duckdb.connect()
    try:
        cursor = connection.execute(query, values)
        headers = [item[0] for item in cursor.description]
        rows = [dict(zip(headers, row)) for row in cursor.fetchall()]
        rejection_report = build_rejection_report(
            connection, path, base_clauses, values, entry_expression,
            drawdown_expression, query, require_robustness,
        )
        diagnostics: dict[str, Any] = {}
        if not rows:
            diagnostic_query = f"""
                WITH source AS (
                  SELECT *, {parameter_key} AS parameter_key,
                         {entry_expression} AS entry_minutes,
                         {drawdown_expression} AS ranking_drawdown
                    FROM read_parquet(?)
                ),
                time_pnl AS (
                  SELECT parameter_key, entry_minutes,
                         max(CASE WHEN lower(CAST({identifier(status)} AS VARCHAR)) = 'succeeded'
                                  THEN CAST({identifier(net_pnl)} AS DOUBLE) END) AS shifted_pnl
                    FROM source
                   WHERE entry_minutes IS NOT NULL
                   GROUP BY parameter_key, entry_minutes
                ),
                base AS (
                  SELECT parameter_key, entry_minutes,
                         max(CAST({identifier(net_pnl)} AS DOUBLE)) AS net_pnl
                    FROM source
                   WHERE {' AND '.join(base_clauses)}
                   GROUP BY parameter_key, entry_minutes
                ),
                robustness_stats AS (
                  SELECT b.*,
                         count(variant.entry_minutes) FILTER (WHERE variant.entry_minutes < b.entry_minutes) AS before_variant_count,
                         count(variant.entry_minutes) FILTER (WHERE variant.entry_minutes > b.entry_minutes) AS after_variant_count,
                         bool_and(coalesce(variant.shifted_pnl >= 0.70 * b.net_pnl, false))
                           FILTER (WHERE variant.entry_minutes < b.entry_minutes) AS before_all_pass,
                         bool_and(coalesce(variant.shifted_pnl >= 0.70 * b.net_pnl, false))
                           FILTER (WHERE variant.entry_minutes > b.entry_minutes) AS after_all_pass
                    FROM base b
                    LEFT JOIN time_pnl variant ON variant.parameter_key = b.parameter_key
                                          AND variant.entry_minutes <> b.entry_minutes
                                          AND variant.entry_minutes BETWEEN b.entry_minutes - {robustness_window_minutes}
                                                                       AND b.entry_minutes + {robustness_window_minutes}
                   GROUP BY ALL
                ),
                all_variants AS (
                  SELECT min(abs(variant.entry_minutes - b.entry_minutes)) AS nearest_variant_minutes,
                         count(*) AS same_config_variant_count
                    FROM base b
                    JOIN time_pnl variant ON variant.parameter_key = b.parameter_key
                                         AND variant.entry_minutes <> b.entry_minutes
                )
                SELECT
                    (SELECT count(*) FROM base) AS base_count,
                    (SELECT same_config_variant_count FROM all_variants) AS same_config_variant_count,
                    (SELECT nearest_variant_minutes FROM all_variants) AS nearest_variant_minutes,
                    count(*) FILTER (WHERE before_variant_count > 0 OR after_variant_count > 0) AS base_with_window_variants,
                    count(*) FILTER (
                        WHERE (before_variant_count > 0 AND coalesce(before_all_pass, false))
                           OR (after_variant_count > 0 AND coalesce(after_all_pass, false))
                    ) AS robustness_pass_count
                  FROM robustness_stats
            """
            diagnostic_row = connection.execute(diagnostic_query, values).fetchone()
            if diagnostic_row:
                diagnostics = {
                    "base_count": int(diagnostic_row[0] or 0),
                    "same_config_variant_count": int(diagnostic_row[1] or 0),
                    "nearest_variant_minutes": int(diagnostic_row[2]) if diagnostic_row[2] is not None else None,
                    "base_with_window_variants": int(diagnostic_row[3] or 0),
                    "robustness_pass_count": int(diagnostic_row[4] or 0),
                    "robustness_window_minutes": robustness_window_minutes,
                }
    except duckdb.Error as exc:
        raise ValueError(f"The sweep could not be ranked: {exc}") from exc
    finally:
        connection.close()

    distribution = {
        "before": int(rows[0]["before_count"]) if rows else 0,
        "after": int(rows[0]["after_count"]) if rows else 0,
        "both": int(rows[0]["both_count"]) if rows else 0,
        "not_confirmed": int(rows[0]["not_confirmed_count"]) if rows else 0,
    }
    result_rows = []
    for row in rows:
        parameters = {column: row.pop(column) for column in parameter_columns}
        # Expose cutoff separately as well as retaining the original parameter.
        # This makes it available to fixed dashboard columns and detail views,
        # even when a sweep producer uses a cutoff alias.
        cutoff_time = next(
            (
                parameters.get(name)
                for name in ("cutoff_time", "cutoff", "cutoff_point", "cutoff_start")
                if parameters.get(name) not in (None, "")
            ),
            None,
        )
        row["cutoff_time"] = cutoff_time
        row.pop("before_count")
        row.pop("after_count")
        row.pop("both_count")
        row.pop("not_confirmed_count")
        components = []
        for item in criteria:
            score = row.pop(f"criterion_score_{item['index']}")
            components.append({
                "column": item["column"],
                "label": item["label"],
                "direction": item["direction"],
                "weight": item["weight"],
                "value": row.pop(f"criterion_value_{item['index']}"),
                "percentile_score": score,
                "weighted_contribution": score * item["weight"] / 100.0,
            })
        result_rows.append({**row, "parameters": parameters, "ranking_components": components})
    return {
        "source_row_count": report["file"]["row_count"],
        "rejection_report": rejection_report,
        "eligible_count": int(rows[0]["eligible_count"]) if rows else 0,
        "top_n": top_n,
        "ranking": [{key: item[key] for key in ("column", "label", "weight", "direction")} for item in criteria],
        "robustness_required": require_robustness,
        "robustness_distribution": distribution,
        "diagnostics": diagnostics if not rows else {},
        "rows": result_rows,
    }

